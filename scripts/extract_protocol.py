#!/usr/bin/env python3
"""Extract DEATH PROTOCOL training plan from xlsx to JSON for seeding."""
import json
import openpyxl

SRC = '/home/z/my-project/upload/death_protocol_training_plan.xlsx'
OUT = '/home/z/my-project/scripts/protocol_data.json'

wb = openpyxl.load_workbook(SRC, data_only=True)

# --- Weekly Protocol sheet -> exercises list ---
ws = wb['Weekly Protocol']
exercises = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        continue
    day, block, name, sets, reps, rest, focus, complete = [str(c).strip() if c is not None else '' for c in row]
    if not name:
        continue
    exercises.append({
        'day': day,
        'block': block,
        'name': name,
        'sets': sets,
        'reps': reps,
        'rest': rest,
        'focus': focus,
        'order': len(exercises),
    })

# --- Overview sheet -> meta (non-empty cells) ---
ov = wb['Overview']
overview_rows = []
for row in ov.iter_rows(values_only=True):
    cells = [str(c).strip() if c is not None else '' for c in row]
    if any(cells):
        overview_rows.append(cells)

data = {'exercises': exercises, 'overview': overview_rows}
with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=1)

days = sorted({e['day'] for e in exercises}, key=lambda d: exercises[[e2['day'] for e2 in exercises].index(d)]['order'])
blocks = sorted({e['block'] for e in exercises})
print('total exercises:', len(exercises))
print('days:', days)
print('blocks:', blocks)
print('overview rows:', len(overview_rows))
